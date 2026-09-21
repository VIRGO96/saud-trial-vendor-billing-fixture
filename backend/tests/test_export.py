import io
from decimal import Decimal
import openpyxl
import pytest
from app.services.comparison import compare_invoices
from app.services.exporter import export_changes_to_csv, export_changes_to_xlsx
from tests.test_comparison import load_sherweb_samples

def test_export_csv_structure():
    items_jul, items_aug = load_sherweb_samples()
    comp = compare_invoices(items_jul, items_aug, "2026-07", "2026-08")
    csv_bytes = export_changes_to_csv(comp, vendor_name="Sherweb")

    assert csv_bytes.startswith(b'\xef\xbb\xbf') # BOM
    text = csv_bytes.decode('utf-8-sig')
    lines = text.strip().split('\r\n')
    if len(lines) == 1:
        lines = text.strip().split('\n')

    # Header + 6 changes
    assert len(lines) == 7
    header = lines[0]
    assert "vendor,period_from,period_to,change_type,account,sku" in header
    assert "Bayview Clinic" in text
    assert "Saltbox Kitchen" in text
    assert "Rosewood Pictures" in text

def test_export_xlsx_structure():
    items_jul, items_aug = load_sherweb_samples()
    comp = compare_invoices(items_jul, items_aug, "2026-07", "2026-08")
    xlsx_bytes = export_changes_to_xlsx(comp, vendor_name="Sherweb")

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Summary" in wb.sheetnames
    assert "Changes" in wb.sheetnames

    # Check Summary Sheet
    ws_summary = wb["Summary"]
    assert "Sherweb Billing Reconciliation Summary" in ws_summary['A1'].value
    assert ws_summary['B3'].value == "2026-07 -> 2026-08"
    assert ws_summary['B7'].value == "Reconciled to the cent"

    # Check Changes Sheet
    ws_changes = wb["Changes"]
    # Row 1 is header, rows 2..7 are the 6 changes
    assert ws_changes.max_row == 7
    header_cells = [ws_changes.cell(row=1, column=col).value for col in range(1, 19)]
    assert "Vendor" in header_cells
    assert "Change Type" in header_cells
    assert "Amount Delta" in header_cells